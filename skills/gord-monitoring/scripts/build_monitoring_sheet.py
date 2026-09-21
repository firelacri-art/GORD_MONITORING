#!/usr/bin/env python3
"""
build_monitoring_sheet.py — лист «Мониторинг публикаций» в стиле GORD из entries.json.

Структура повторяет таблицу агентства: логотип, заголовок, шапка
Название · Ссылка · Формат · AVE (стоимость) · OTS (охват) · Дата · Комментарий; строки сгруппированы
по месяцам, внутри месяца — подсекции сми / тг / инст / тикток; в конце ИТОГО.
Ссылка вставляется в ячейку как обычный URL (не «ссылка» с гиперссылкой) — так просил партнёр.

    python3 build_monitoring_sheet.py <slug|project.json> [out.xlsx] [--csv out.csv] [--since 2026-04-01]
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import sys
from collections import OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "vendor"))
from monitoring_lib import ASSETS, CHANNELS, State, load_project, month_label  # noqa: E402

import build_gord_table as gt  # noqa: E402  (vendored генератор gord-tables)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

COLUMNS = [  # (заголовок, ширина, выравнивание)
    ("Название", 26, "left"),
    ("Ссылка", 58, "left"),
    ("Формат", 24, "left"),
    ("AVE (стоимость)", 18, "center"),
    ("OTS (охват)", 18, "center"),
    ("Дата", 13, "center"),
    ("Комментарий", 30, "left"),
]
CHANNEL_TITLES = {"сми": "СМИ", "тг": "Telegram", "инст": "Instagram", "тикток": "TikTok"}


def group(entries: list[dict]) -> "OrderedDict[str, OrderedDict[str, list[dict]]]":
    out: OrderedDict = OrderedDict()
    for e in sorted(entries, key=lambda x: (x.get("date") or "", x.get("channel", ""), x.get("name", ""))):
        m = month_label(e.get("date", ""))
        out.setdefault(m, OrderedDict())
        out[m].setdefault(e.get("channel", "сми"), []).append(e)
    return out


def build(entries: list[dict], out: Path, project: str, logo: Path) -> Path:
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Мониторинг публикаций")
    gt.page_setup(ws)
    ws.column_dimensions["A"].width = gt.COL_MARGIN
    first, last = 2, 1 + len(COLUMNS)
    for i, (_, w, _) in enumerate(COLUMNS):
        ws.column_dimensions[get_column_letter(first + i)].width = w
    ws.column_dimensions[get_column_letter(last + 1)].width = gt.COL_MARGIN

    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = gt.ROW_LOGO
    gt.add_logo(ws, logo, cell_row=2)
    red = Side(style="medium", color="FF" + gt.RED)
    for c in range(first, last + 1):
        ws.cell(row=2, column=c).border = Border(bottom=red)
    ws.row_dimensions[3].height = gt.ROW_TITLE
    ws.merge_cells(start_row=3, start_column=first, end_row=3, end_column=last)
    t = ws.cell(row=3, column=first, value=f"Мониторинг публикаций — {project}")
    t.font = gt.font(gt.SIZE_TITLE, bold=True); t.alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[4].height = gt.ROW_HEADER
    for i, (h, _, _) in enumerate(COLUMNS):
        c = ws.cell(row=4, column=first + i, value=h)
        c.fill = gt.fill(gt.GREY_HEADER); c.border = gt.box(); c.font = gt.font(gt.SIZE_HEADER, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r = 5
    ave_cells, ots_cells = [], []

    def section(label: str, r: int, sub: bool = False) -> None:
        ws.row_dimensions[r].height = gt.ROW_BODY
        ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
        for c in range(first, last + 1):
            ws.cell(row=r, column=c).border = gt.box()
        cell = ws.cell(row=r, column=first, value=label)
        cell.font = gt.font(gt.SIZE_BODY, bold=True)
        cell.alignment = Alignment(horizontal="left" if sub else "center", vertical="center", indent=1 if sub else 0)
        if not sub:
            for c in range(first, last + 1):
                b = ws.cell(row=r, column=c).border
                ws.cell(row=r, column=c).border = Border(left=b.left, right=b.right, top=b.top, bottom=red)
        elif sub:
            for c in range(first, last + 1):
                ws.cell(row=r, column=c).fill = gt.fill(gt.GREY_ZEBRA)

    for month, channels in group(entries).items():
        section(month, r); r += 1
        for ch in CHANNELS:
            items = channels.get(ch)
            if not items:
                continue
            if len(channels) > 1:
                section(CHANNEL_TITLES[ch], r, sub=True); r += 1
            for e in items:
                ws.row_dimensions[r].height = gt.ROW_BODY
                vals = [e.get("name", ""), e.get("url", ""), e.get("format", ""), e.get("ave"), e.get("ots"),
                        dt.date.fromisoformat(e["date"][:10]) if e.get("date") else None, e.get("note", "")]
                for i, ((_, _, al), v) in enumerate(zip(COLUMNS, vals)):
                    c = ws.cell(row=r, column=first + i, value=v)
                    c.font = gt.font(gt.SIZE_BODY); c.border = gt.box()
                    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=(i not in (1, 6)))
                    if i in (3, 4):
                        c.number_format = gt.NUM_FMT
                    if i == 5:
                        c.number_format = gt.DATE_FMT
                ave_cells.append(f"{get_column_letter(first + 3)}{r}"); ots_cells.append(f"{get_column_letter(first + 4)}{r}")
                r += 1
        r += 1  # воздух между месяцами

    # ИТОГО
    ws.row_dimensions[r].height = gt.ROW_BODY
    ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=first + 2)
    for c in range(first, last + 1):
        cell = ws.cell(row=r, column=c); cell.fill = gt.fill(gt.GREY_TOTAL); cell.border = gt.box(); cell.font = gt.font(gt.SIZE_BODY, bold=True)
    ws.cell(row=r, column=first, value=f"ИТОГО: {len(entries)} публикаций")
    if ave_cells:
        ca = ws.cell(row=r, column=first + 3, value="=" + "+".join(ave_cells) if len(ave_cells) < 200 else f"=SUM({get_column_letter(first + 3)}5:{get_column_letter(first + 3)}{r - 1})")
        co = ws.cell(row=r, column=first + 4, value="=" + "+".join(ots_cells) if len(ots_cells) < 200 else f"=SUM({get_column_letter(first + 4)}5:{get_column_letter(first + 4)}{r - 1})")
        for c in (ca, co):
            c.number_format = gt.NUM_FMT; c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = ws.cell(row=5, column=first)
    wb.properties.creator = "GORD Agency"; wb.properties.title = f"Мониторинг публикаций — {project}"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def write_csv(entries: list[dict], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([h for h, _, _ in COLUMNS] + ["Канал", "Заголовок"])
        for e in sorted(entries, key=lambda x: x.get("date", "")):
            w.writerow([e.get("name"), e.get("url"), e.get("format"), e.get("ave"), e.get("ots"), e.get("date"), e.get("note", ""), e.get("channel"), e.get("title")])


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("project"); ap.add_argument("out", nargs="?")
    ap.add_argument("--csv", help="дополнительно выгрузить CSV (удобно вставлять в Google-таблицу)")
    ap.add_argument("--since", help="только публикации с этой даты (ISO)")
    ap.add_argument("--logo", help="свой логотип PNG")
    a = ap.parse_args(argv)
    slug, cfg = load_project(a.project)
    st = State(slug)
    entries = [e for e in st.entries if not a.since or (e.get("date") or "") >= a.since]
    if not entries:
        sys.exit("Нет подтверждённых публикаций (entries.json пуст) — сначала add_entry.py")
    out = Path(a.out or (st.dir / f"Мониторинг публикаций — {cfg['project']}.xlsx"))
    build(entries, out, cfg["project"], Path(a.logo or ASSETS / "gord-logo.png"))
    if a.csv:
        write_csv(entries, Path(a.csv))
    ave = sum(e.get("ave") or 0 for e in entries); ots = sum(e.get("ots") or 0 for e in entries)
    print(f"OK → {out}\nПубликаций: {len(entries)}; AVE: {ave:,}; OTS: {ots:,}".replace(",", " "))


if __name__ == "__main__":
    main()
